#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
将 doc/编码分类规则以及编码属性说明.xlsx 的「生产物资-2023」工作表
转换为系统可读的配置：
  - src/config/classification.json   完整分类树（大类/中类/小类 + 库房 + 编码pattern）
  - src/config/classification-map.json  6位小类编码 -> 分类定位（供 FNUMBER 前缀归类）

用法： python scripts/gen_classification.py
依赖： openpyxl
说明： 工作表中 大类(A/B)、中类(C/D) 为合并单元格，仅首行有值，本脚本做前向填充。
"""
import json
import os
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "doc", "编码分类规则以及编码属性说明.xlsx")
CFG_DIR = os.path.join(ROOT, "src", "config")
os.makedirs(CFG_DIR, exist_ok=True)

wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb["生产物资-2023"]

categories = []          # 大类列表
cur_big = None           # 当前大类 dict
cur_mid = None           # 当前中类 dict
sub_map = {}             # 6位小类编码 -> 定位信息

# 前向填充的临时变量
last_big_code = last_big_name = None
last_mid_code = last_mid_name = None


def flush_mid():
    global cur_mid, cur_big
    if cur_big is not None and cur_mid is not None:
        cur_big["mid"].append(cur_mid)
        cur_mid = None


def flush_big():
    global cur_big
    flush_mid()
    if cur_big is not None:
        categories.append(cur_big)
        cur_big = None


for r in range(3, ws.max_row + 1):
    b = ws.cell(row=r, column=2).value   # 大类编码
    a = ws.cell(row=r, column=1).value   # 大类名称
    d = ws.cell(row=r, column=4).value   # 中类编码
    c = ws.cell(row=r, column=3).value   # 中类名称
    f = ws.cell(row=r, column=6).value   # 小类编码
    e = ws.cell(row=r, column=5).value   # 小类名称
    i = ws.cell(row=r, column=9).value   # 编码 pattern
    j = ws.cell(row=r, column=10).value  # 单位
    k = ws.cell(row=r, column=11).value  # 库房名称

    if b:
        last_big_code, last_big_name = str(b).strip(), (str(a).strip() if a else "")
    if d:
        last_mid_code, last_mid_name = str(d).strip(), (str(c).strip() if c else "")
    if not f:
        continue  # 非小类行，跳过（空行/只填了大类的行）

    f = str(f).strip()
    e = str(e).strip() if e else ""
    i = str(i).strip() if i else ""
    j = str(j).strip() if j else ""
    k = str(k).strip() if k else ""

    # 大类切换
    if last_big_code != (cur_big["code"] if cur_big else None):
        flush_big()
        cur_big = {"code": last_big_code, "name": last_big_name, "mid": []}
        cur_mid = None
    # 中类切换
    if last_mid_code != (cur_mid["code"] if cur_mid else None):
        flush_mid()
        cur_mid = {"code": last_mid_code, "name": last_mid_name, "sub": []}

    cur_mid["sub"].append({
        "code": f,
        "name": e,
        "pattern": i,
        "unit": j,
        "wh": k,
    })

    sub_map[f] = {
        "big": last_big_code,
        "bigName": last_big_name,
        "mid": last_mid_code,
        "midName": last_mid_name,
        "sub": f,
        "subName": e,
        "wh": k,
    }

flush_big()

classification = {
    "meta": {
        "source": "doc/编码分类规则以及编码属性说明.xlsx / 生产物资-2023",
        "encodingRule": "物料编码 = 小类编码(6位 = 大类2位+中类2位+小类序号2位) + 5位流水码 = 11位",
        "generatedBy": "scripts/gen_classification.py",
        "bigCount": len(categories),
        "subCount": len(sub_map),
    },
    "categories": categories,
}

with open(os.path.join(CFG_DIR, "classification.json"), "w", encoding="utf-8") as fh:
    json.dump(classification, fh, ensure_ascii=False, indent=2)

with open(os.path.join(CFG_DIR, "classification-map.json"), "w", encoding="utf-8") as fh:
    json.dump(sub_map, fh, ensure_ascii=False, indent=2)

print(f"OK  大类={len(categories)}  小类={len(sub_map)}")
print(f"  -> {os.path.join(CFG_DIR,'classification.json')}")
print(f"  -> {os.path.join(CFG_DIR,'classification-map.json')}")
